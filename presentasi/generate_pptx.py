"""
RoRo Smart Rollator — Generator Presentasi PPTX
Jalankan: python generate_pptx.py
Output:   RoRo_Presentasi.pptx  (di folder ini)

Dependensi:
    pip install python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import sys

# ---------------------------------------------------------------------------
# PALET WARNA
# ---------------------------------------------------------------------------
C_NAVY    = RGBColor(0x0D, 0x1B, 0x2A)   # latar gelap
C_BLUE    = RGBColor(0x1B, 0x63, 0xB0)   # aksen biru
C_CYAN    = RGBColor(0x00, 0xC2, 0xCB)   # aksen cyan / highlight
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT   = RGBColor(0xE8, 0xF4, 0xFD)   # bg slide konten
C_GRAY    = RGBColor(0x55, 0x65, 0x78)   # teks sekunder
C_GREEN   = RGBColor(0x2E, 0xCC, 0x71)   # aksen sukses/aktif
C_ORANGE  = RGBColor(0xE6, 0x7E, 0x22)   # aksen peringatan

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# ---------------------------------------------------------------------------
# HELPER
# ---------------------------------------------------------------------------
def prs_new():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # completely blank
    return prs.slides.add_slide(layout)


def bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def box(slide, left, top, width, height,
        fill_color=None, line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE
        left, top, width, height
    )
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def txt(slide, text, left, top, width, height,
        font_size=Pt(18), bold=False, color=C_WHITE,
        align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = slide.shapes.add_textbox(left, top, width, height)
    txb.word_wrap = wrap
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = font_size
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txb


def add_bullet_box(slide, title, bullets, left, top, width, height,
                   title_color=C_CYAN, bullet_color=C_NAVY,
                   bg_color=C_WHITE, accent=C_BLUE):
    """Kotak dengan judul dan bullet list."""
    # shadow box
    box(slide, left + Inches(0.05), top + Inches(0.05), width, height,
        fill_color=RGBColor(0xCC, 0xDD, 0xEE))
    # main box
    b = box(slide, left, top, width, height, fill_color=bg_color,
            line_color=accent, line_width=Pt(1.5))
    # accent bar kiri
    box(slide, left, top, Inches(0.12), height, fill_color=accent)
    # title
    txt(slide, title,
        left + Inches(0.2), top + Inches(0.08),
        width - Inches(0.25), Inches(0.45),
        font_size=Pt(14), bold=True, color=title_color, align=PP_ALIGN.LEFT)
    # bullets
    bullet_top = top + Inches(0.55)
    bullet_h   = (height - Inches(0.6)) / max(len(bullets), 1)
    for b_text in bullets:
        txt(slide, f"  \u25CF  {b_text}",
            left + Inches(0.2), bullet_top,
            width - Inches(0.3), bullet_h,
            font_size=Pt(11.5), color=bullet_color, align=PP_ALIGN.LEFT)
        bullet_top += bullet_h


def accent_line(slide, left, top, width, color=C_CYAN, thick=Pt(3)):
    ln = slide.shapes.add_shape(1, left, top, width, thick)
    ln.fill.solid()
    ln.fill.fore_color.rgb = color
    ln.line.fill.background()


def section_header(slide, number, title, subtitle=""):
    """Header section atas slide konten."""
    # bar atas
    box(slide, 0, 0, SLIDE_W, Inches(1.1), fill_color=C_NAVY)
    # nomor badge
    box(slide, Inches(0.3), Inches(0.18), Inches(0.65), Inches(0.65),
        fill_color=C_CYAN)
    txt(slide, str(number),
        Inches(0.3), Inches(0.18), Inches(0.65), Inches(0.65),
        font_size=Pt(22), bold=True, color=C_NAVY, align=PP_ALIGN.CENTER)
    # judul
    txt(slide, title,
        Inches(1.1), Inches(0.15), Inches(9), Inches(0.5),
        font_size=Pt(24), bold=True, color=C_WHITE)
    if subtitle:
        txt(slide, subtitle,
            Inches(1.1), Inches(0.62), Inches(9), Inches(0.35),
            font_size=Pt(13), color=C_CYAN, italic=True)
    # garis aksen
    accent_line(slide, 0, Inches(1.1), SLIDE_W, color=C_CYAN, thick=Pt(3))


def footer(slide, page_num, total):
    box(slide, 0, Inches(7.1), SLIDE_W, Inches(0.4), fill_color=C_NAVY)
    txt(slide, "RoRo — Smart Rollator   |   Firmware ESP32",
        Inches(0.3), Inches(7.12), Inches(9), Inches(0.3),
        font_size=Pt(9), color=C_GRAY)
    txt(slide, f"{page_num} / {total}",
        Inches(12.5), Inches(7.12), Inches(0.6), Inches(0.3),
        font_size=Pt(9), color=C_CYAN, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# SLIDE BUILDER
# ---------------------------------------------------------------------------

def slide_cover(prs):
    s = blank_slide(prs)
    bg(s, C_NAVY)

    # gradient overlay kiri ke kanan (simulasi pakai 2 kotak)
    ov = box(s, 0, 0, Inches(8), SLIDE_H, fill_color=RGBColor(0x0A, 0x14, 0x22))
    box(s, Inches(8), 0, Inches(5.33), SLIDE_H, fill_color=C_BLUE)

    # garis dekorasi
    accent_line(s, 0,           Inches(1.8), SLIDE_W, C_CYAN, Pt(3))
    accent_line(s, 0,           Inches(1.95), SLIDE_W, RGBColor(0x1B,0x63,0xB0), Pt(1))
    accent_line(s, 0,           Inches(5.5), Inches(7), C_CYAN, Pt(2))

    # badge produk
    box(s, Inches(0.5), Inches(0.5), Inches(1.8), Inches(0.6),
        fill_color=C_CYAN)
    txt(s, "TUGAS AKHIR",
        Inches(0.5), Inches(0.5), Inches(1.8), Inches(0.6),
        font_size=Pt(11), bold=True, color=C_NAVY, align=PP_ALIGN.CENTER)

    # judul utama
    txt(s, "RoRo",
        Inches(0.5), Inches(2.0), Inches(7.5), Inches(1.5),
        font_size=Pt(72), bold=True, color=C_WHITE, align=PP_ALIGN.LEFT)

    # sub-judul
    txt(s, "Smart Robot Rollator",
        Inches(0.5), Inches(3.4), Inches(7.5), Inches(0.7),
        font_size=Pt(28), bold=False, color=C_CYAN, align=PP_ALIGN.LEFT)

    txt(s, "Tongkat Jalan Pintar Berbasis ESP32\ndengan Kendali Motor, Deteksi IMU, SOS & Monitoring Realtime",
        Inches(0.5), Inches(4.05), Inches(7.5), Inches(1.0),
        font_size=Pt(14), color=RGBColor(0xB0, 0xC8, 0xE0), align=PP_ALIGN.LEFT)

    # info bawah
    accent_line(s, Inches(0.5), Inches(5.3), Inches(3.5), C_CYAN, Pt(1.5))
    txt(s, "Tim Pengembang  |  2025",
        Inches(0.5), Inches(5.45), Inches(4), Inches(0.35),
        font_size=Pt(11), color=RGBColor(0xB0, 0xC8, 0xE0))
    txt(s, "Platform: ESP32  \u2022  Framework: Arduino / PlatformIO\nBahasa: C++  \u2022  Real-time FreeRTOS  \u2022  Cloud: Firebase + HiveMQ",
        Inches(0.5), Inches(5.8), Inches(6), Inches(0.7),
        font_size=Pt(10), color=C_GRAY)

    # ikon / grafis sisi kanan (kotak representasi)
    for i, (lbl, clr) in enumerate([
        ("ESP32", C_CYAN),
        ("IMU", C_GREEN),
        ("MQTT", C_ORANGE),
        ("Firebase", RGBColor(0xFF, 0xCA, 0x28)),
    ]):
        bx_l = Inches(8.6) + (i % 2) * Inches(2.1)
        bx_t = Inches(1.5) + (i // 2) * Inches(1.8)
        box(s, bx_l, bx_t, Inches(1.8), Inches(1.4), fill_color=clr)
        txt(s, lbl, bx_l, bx_t + Inches(0.45), Inches(1.8), Inches(0.5),
            font_size=Pt(16), bold=True, color=C_NAVY, align=PP_ALIGN.CENTER)

    txt(s, "Firmware ESP32",
        Inches(8.4), Inches(4.8), Inches(4.5), Inches(0.4),
        font_size=Pt(11), color=C_GRAY, align=PP_ALIGN.CENTER)


def slide_latar_belakang(prs):
    s = blank_slide(prs)
    bg(s, C_LIGHT)
    section_header(s, "01", "Latar Belakang", "Mengapa RoRo dibangun?")
    footer(s, 2, 10)

    # 3 problem card
    problems = [
        ("Populasi Lansia Meningkat",
         "Jumlah lanjut usia di Indonesia terus bertambah.\nRisiko jatuh tinggi — butuh alat bantu yang lebih cerdas."),
        ("Rollator Konvensional Pasif",
         "Tongkat jalan tradisional tidak memberi umpan balik.\nTidak ada monitoring kondisi pengguna secara realtime."),
        ("Keselamatan & Respons Darurat",
         "Tidak ada mekanisme SOS otomatis.\nKeluarga / perawat tidak bisa memantau jarak jauh."),
    ]
    colors = [C_BLUE, C_ORANGE, RGBColor(0xC0, 0x39, 0x2B)]
    for i, (title, desc) in enumerate(problems):
        left = Inches(0.4) + i * Inches(4.2)
        box(s, left, Inches(1.35), Inches(3.9), Inches(3.2), fill_color=colors[i])
        txt(s, title, left + Inches(0.15), Inches(1.5), Inches(3.6), Inches(0.55),
            font_size=Pt(15), bold=True, color=C_WHITE)
        accent_line(s, left + Inches(0.15), Inches(2.05), Inches(3.6),
                    RGBColor(0xFF,0xFF,0xFF), Pt(1))
        txt(s, desc, left + Inches(0.15), Inches(2.15), Inches(3.6), Inches(2.0),
            font_size=Pt(12), color=C_WHITE)

    # solusi
    box(s, Inches(0.4), Inches(4.75), Inches(12.5), Inches(1.5),
        fill_color=C_NAVY)
    txt(s, "\u2714  Solusi: RoRo — Smart Rollator",
        Inches(0.6), Inches(4.82), Inches(5), Inches(0.45),
        font_size=Pt(16), bold=True, color=C_CYAN)
    txt(s,
        "Rollator bermotor dengan kendali kecepatan adaptif  \u2022  Deteksi gerakan IMU (berjalan vs diam)  \u2022  Tombol SOS + notifikasi realtime\n"
        "Dashboard lokal via HTTP  \u2022  History tersimpan di Firebase Firestore  \u2022  Push instan ke Flutter app via MQTT HiveMQ",
        Inches(0.6), Inches(5.3), Inches(12.1), Inches(0.8),
        font_size=Pt(11), color=C_WHITE)


def slide_arsitektur(prs):
    s = blank_slide(prs)
    bg(s, C_LIGHT)
    section_header(s, "02", "Arsitektur Sistem", "Dua domain eksekusi — realtime & networking terpisah")
    footer(s, 3, 10)

    # domain kiri: FreeRTOS
    box(s, Inches(0.4), Inches(1.3), Inches(5.5), Inches(5.5),
        fill_color=C_NAVY, line_color=C_CYAN, line_width=Pt(2))
    txt(s, "Core 1 — FreeRTOS Tasks (Realtime)",
        Inches(0.55), Inches(1.38), Inches(5.2), Inches(0.45),
        font_size=Pt(13), bold=True, color=C_CYAN)
    accent_line(s, Inches(0.4), Inches(1.83), Inches(5.5), C_CYAN, Pt(2))

    tasks = [
        ("TaskSafety", "Prio 3 \u2022 20 ms", "SOS button, buzzer state machine, speed beep"),
        ("TaskMotor",  "Prio 2 \u2022 10 ms", "BTS7960 PWM, potentiometer, gas button + latch"),
        ("TaskIMU",    "Prio 2 \u2022 120 ms","MPU6050 I2C, motionScore EMA, deteksi berjalan/diam"),
    ]
    colors_t = [C_ORANGE, C_BLUE, C_GREEN]
    for i, (name, prio, desc) in enumerate(tasks):
        top = Inches(2.0) + i * Inches(1.5)
        box(s, Inches(0.55), top, Inches(5.1), Inches(1.3),
            fill_color=colors_t[i])
        txt(s, name, Inches(0.7), top + Inches(0.05), Inches(3), Inches(0.45),
            font_size=Pt(14), bold=True, color=C_WHITE)
        txt(s, prio, Inches(3.7), top + Inches(0.08), Inches(1.8), Inches(0.35),
            font_size=Pt(10), color=C_WHITE, align=PP_ALIGN.RIGHT)
        txt(s, desc, Inches(0.7), top + Inches(0.5), Inches(4.8), Inches(0.6),
            font_size=Pt(11), color=C_WHITE)

    # panah tengah
    txt(s, "Queue*()\nnon-blocking",
        Inches(6.0), Inches(3.3), Inches(1.3), Inches(0.9),
        font_size=Pt(10), bold=True, color=C_BLUE, align=PP_ALIGN.CENTER)
    txt(s, "\u27A1", Inches(6.1), Inches(3.6), Inches(1.0), Inches(0.5),
        font_size=Pt(30), color=C_CYAN, align=PP_ALIGN.CENTER)

    # domain kanan: Arduino loop
    box(s, Inches(7.4), Inches(1.3), Inches(5.5), Inches(5.5),
        fill_color=RGBColor(0x1A, 0x2A, 0x40), line_color=C_BLUE, line_width=Pt(2))
    txt(s, "Arduino loop() — Networking (Best-Effort)",
        Inches(7.55), Inches(1.38), Inches(5.2), Inches(0.45),
        font_size=Pt(13), bold=True, color=C_BLUE)
    accent_line(s, Inches(7.4), Inches(1.83), Inches(5.5), C_BLUE, Pt(2))

    loops = [
        ("wifiManagerLoop",    "WiFi AP/STA state machine, mDNS rorro.local"),
        ("firebaseManagerLoop","Firestore REST HTTPS — history & storage"),
        ("mqttManagerLoop",    "HiveMQ TLS 8883 — push realtime ke Flutter"),
        ("webServerLoop",      "HTTP API port 80 + dashboard.html"),
    ]
    lp_colors = [C_CYAN, RGBColor(0xFF,0xCA,0x28), C_ORANGE, C_GREEN]
    for i, (name, desc) in enumerate(loops):
        top = Inches(2.05) + i * Inches(1.15)
        box(s, Inches(7.55), top, Inches(5.1), Inches(1.0),
            fill_color=RGBColor(0x0D,0x22,0x36))
        box(s, Inches(7.55), top, Inches(0.12), Inches(1.0),
            fill_color=lp_colors[i])
        txt(s, name, Inches(7.75), top + Inches(0.05), Inches(4.7), Inches(0.38),
            font_size=Pt(12), bold=True, color=lp_colors[i])
        txt(s, desc, Inches(7.75), top + Inches(0.42), Inches(4.7), Inches(0.45),
            font_size=Pt(10), color=C_WHITE)

    # shared state
    box(s, Inches(0.4), Inches(6.9), Inches(12.5), Inches(0.4),
        fill_color=RGBColor(0x1B, 0x3A, 0x5C))
    txt(s, "Shared State: RobotData struct  \u2022  portMUX_TYPE critical sections  \u2022  volatile gActiveMotorSpeed / gSpeedLevelChanged",
        Inches(0.55), Inches(6.92), Inches(12.1), Inches(0.35),
        font_size=Pt(10), color=C_CYAN, align=PP_ALIGN.CENTER)


def slide_fitur_utama(prs):
    s = blank_slide(prs)
    bg(s, C_LIGHT)
    section_header(s, "03", "Fitur Utama", "Lima modul firmware RoRo")
    footer(s, 4, 10)

    features = [
        ("Kendali Motor DC",
         C_BLUE,
         ["BTS7960 H-bridge x2 (kiri & kanan)",
          "PWM 5 kHz, 8-bit, 0-250 level",
          "9 level kecepatan via potensiometer",
          "Tombol gas + latch triple-tap (3x ketuk = terkunci)",
          "Konfirmasi beep saat ganti level kecepatan"]),
        ("Deteksi Gerak IMU",
         C_GREEN,
         ["MPU6050 6-axis I2C (GPIO 21/22)",
          "motionScore = accelDelta \u00d7 3.5 + gyroMag / 45",
          "EMA smoothing \u03b1=0.25 @ 120 ms",
          "Histerisis: ON \u22650.32, OFF \u22640.18",
          "Status: berjalan / diam"]),
        ("SOS & Keselamatan",
         C_ORANGE,
         ["Tombol SOS GPIO 4 (active-LOW)",
          "Buzzer GPIO 5 (toggle 250 ms) saat SOS",
          "Notifikasi Firebase + MQTT instan",
          "TaskSafety prioritas tertinggi (prio 3)",
          "Motor stop sebelum aksi jaringan"]),
        ("Dashboard Lokal & API",
         C_CYAN,
         ["HTTP API port 80 (GET/POST)",
          "GET /api/status — full state JSON",
          "POST /api/motor/test — uji motor manual",
          "POST /api/sos/test — simulasi SOS",
          "dashboard.html polling tiap 1 detik"]),
        ("Cloud & Realtime Push",
         RGBColor(0x8E, 0x44, 0xAD),
         ["Firebase Firestore — history & storage",
          "MQTT HiveMQ TLS 8883 — live push",
          "Deferred-queue: tasks tidak blocking",
          "Retained message: subscriber langsung dapat state",
          "Last Will: status offline otomatis"]),
    ]

    col_w = Inches(2.52)
    for i, (title, color, bullets) in enumerate(features):
        left = Inches(0.2) + i * Inches(2.6)
        box(s, left, Inches(1.3), col_w, Inches(5.6), fill_color=color)
        txt(s, title, left + Inches(0.1), Inches(1.38), col_w - Inches(0.2), Inches(0.6),
            font_size=Pt(13), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        accent_line(s, left + Inches(0.1), Inches(1.98), col_w - Inches(0.2),
                    C_WHITE, Pt(1))
        b_top = Inches(2.1)
        for blt in bullets:
            txt(s, f"\u25B6  {blt}",
                left + Inches(0.1), b_top, col_w - Inches(0.15), Inches(0.8),
                font_size=Pt(10), color=C_WHITE)
            b_top += Inches(0.85)


def slide_tech_stack(prs):
    s = blank_slide(prs)
    bg(s, C_NAVY)
    section_header(s, "04", "Tech Stack", "Hardware, firmware, cloud, dan mobile")
    footer(s, 5, 10)

    categories = [
        ("Hardware", C_CYAN, [
            "ESP32 DevKit (Xtensa LX6, 240 MHz, dual-core)",
            "MPU6050 — Accelerometer + Gyroscope 6-axis",
            "BTS7960 H-Bridge Driver x2 (motor kiri & kanan)",
            "DC Motor x2 (rollator kiri & kanan)",
            "Potensiometer 10K (GPIO 34, ADC 12-bit)",
            "Buzzer aktif-LOW (GPIO 5)",
            "Tombol SOS + Gas (INPUT_PULLUP)",
        ]),
        ("Firmware & Framework", C_GREEN, [
            "PlatformIO + Arduino Framework (C++)",
            "FreeRTOS — xTaskCreatePinnedToCore (Core 1)",
            "ArduinoJson v7 — serialisasi JSON",
            "PubSubClient — MQTT client",
            "WiFiClientSecure — TLS tanpa verifikasi sertifikat",
            "WebServer — HTTP sinkron port 80",
            "Preferences — penyimpanan NVS (WiFi credentials)",
        ]),
        ("Cloud & Infrastruktur", C_ORANGE, [
            "Firebase Firestore — REST API (documents:commit)",
            "Google Identity Toolkit — anonymous auth + token",
            "HiveMQ Cloud — MQTT broker TLS port 8883",
            "NTP: pool.ntp.org / time.google.com",
            "mDNS: rorro.local",
        ]),
        ("Mobile & Monitoring", RGBColor(0x8E,0x44,0xAD), [
            "Flutter app (dikembangkan tim mobile)",
            "Subscribe MQTT: rollators/<id>/#",
            "History dari Firestore subcollections",
            "dashboard.html — control panel lokal",
            "Tools: imu_tune.py, mqtt_test.py (Python)",
        ]),
    ]

    col_w = Inches(3.15)
    col_h = Inches(5.5)
    for i, (cat, color, items) in enumerate(categories):
        left = Inches(0.22) + i * Inches(3.25)
        box(s, left, Inches(1.3), col_w, col_h,
            fill_color=RGBColor(0x0D,0x22,0x36),
            line_color=color, line_width=Pt(2))
        # header bar
        box(s, left, Inches(1.3), col_w, Inches(0.55), fill_color=color)
        txt(s, cat, left, Inches(1.33), col_w, Inches(0.45),
            font_size=Pt(14), bold=True, color=C_NAVY, align=PP_ALIGN.CENTER)
        b_top = Inches(2.0)
        for item in items:
            txt(s, f"\u2022  {item}", left + Inches(0.1), b_top,
                col_w - Inches(0.15), Inches(0.6),
                font_size=Pt(10.5), color=C_WHITE)
            b_top += Inches(0.65)


def slide_alur_kerja(prs):
    s = blank_slide(prs)
    bg(s, C_LIGHT)
    section_header(s, "05", "Alur Kerja Sistem", "Dari sensor ke cloud — deferred-queue pattern")
    footer(s, 6, 10)

    # flow boxes: sensor → task → queue → loop → cloud
    steps = [
        ("Sensor / Input", C_NAVY,
         ["Potensiometer (ADC)", "Tombol Gas (GPIO 16)", "Tombol SOS (GPIO 4)", "MPU6050 (I2C)"]),
        ("FreeRTOS Task\n(Core 1 — Realtime)", C_BLUE,
         ["TaskMotor (10 ms)", "TaskSafety (20 ms)", "TaskIMU (120 ms)", "portMUX critical section"]),
        ("Queue*()\nNon-Blocking", C_CYAN,
         ["firebaseManagerQueue*()", "mqttManagerQueue*()", "Hanya set flag + cache value", "Tidak blocking"]),
        ("Arduino loop()\n(Best-Effort)", C_GREEN,
         ["wifiManagerLoop()", "firebaseManagerLoop()", "mqttManagerLoop()", "webServerLoop()"]),
        ("Cloud / Output", C_ORANGE,
         ["Firebase Firestore", "HiveMQ MQTT", "Flutter App", "dashboard.html"]),
    ]

    box_w = Inches(2.35)
    box_h = Inches(4.8)
    for i, (title, color, items) in enumerate(steps):
        left = Inches(0.2) + i * Inches(2.6)
        box(s, left, Inches(1.35), box_w, box_h, fill_color=color)
        txt(s, title, left + Inches(0.08), Inches(1.42), box_w - Inches(0.1), Inches(0.75),
            font_size=Pt(12), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        accent_line(s, left + Inches(0.1), Inches(2.18),
                    box_w - Inches(0.2), C_WHITE, Pt(1))
        b_top = Inches(2.3)
        for item in items:
            txt(s, f"\u25CF  {item}", left + Inches(0.1), b_top,
                box_w - Inches(0.15), Inches(0.55),
                font_size=Pt(10), color=C_WHITE)
            b_top += Inches(0.58)

        # panah antar box
        if i < len(steps) - 1:
            txt(s, "\u27A1",
                left + box_w + Inches(0.05), Inches(3.5),
                Inches(0.25), Inches(0.5),
                font_size=Pt(18), color=C_NAVY, align=PP_ALIGN.CENTER)

    # catatan retry
    box(s, Inches(0.2), Inches(6.25), Inches(12.9), Inches(0.55),
        fill_color=RGBColor(0xFD, 0xF2, 0xE4))
    txt(s, "\u26A0  Retry logic: jika Firestore/MQTT gagal, flag pending=true → coba lagi 5 detik kemudian. "
           "Tasks tidak pernah blocking karena network.",
        Inches(0.35), Inches(6.3), Inches(12.5), Inches(0.45),
        font_size=Pt(10.5), color=C_ORANGE)


def slide_api_dashboard(prs):
    s = blank_slide(prs)
    bg(s, C_LIGHT)
    section_header(s, "06", "API & Dashboard Lokal", "HTTP port 80 — akses via rorro.local")
    footer(s, 7, 10)

    # tabel endpoint kiri
    endpoints = [
        ("GET",  "/api/status",             "Full state: WiFi, gas, SOS, IMU, speeds, uptime"),
        ("GET",  "/api/imu/status",          "Field IMU saja: status, pitch, roll, motionScore"),
        ("POST", "/api/wifi",               "Simpan SSID+password, restart ke STA mode"),
        ("POST", "/api/provisioning/start", "Paksa AP mode, restart untuk re-provisioning"),
        ("POST", "/api/motor/test",         "Uji motor manual: {left, right, duration_ms}"),
        ("POST", "/api/sos/test",           "Simulasi SOS: {clear: bool} → queue ke Firebase"),
    ]
    method_colors = {
        "GET":  C_GREEN,
        "POST": C_ORANGE,
    }
    box(s, Inches(0.3), Inches(1.3), Inches(7.5), Inches(5.5),
        fill_color=C_WHITE, line_color=C_BLUE, line_width=Pt(1.5))
    txt(s, "HTTP Endpoints", Inches(0.4), Inches(1.38), Inches(3), Inches(0.4),
        font_size=Pt(14), bold=True, color=C_BLUE)
    accent_line(s, Inches(0.3), Inches(1.82), Inches(7.5), C_BLUE, Pt(1.5))

    for i, (method, path, desc) in enumerate(endpoints):
        row_top = Inches(2.0) + i * Inches(0.75)
        box(s, Inches(0.35), row_top, Inches(0.75), Inches(0.5),
            fill_color=method_colors[method])
        txt(s, method, Inches(0.35), row_top + Inches(0.06), Inches(0.75), Inches(0.38),
            font_size=Pt(10), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        txt(s, path, Inches(1.15), row_top + Inches(0.05), Inches(2.8), Inches(0.38),
            font_size=Pt(10.5), bold=True, color=C_NAVY)
        txt(s, desc, Inches(1.15), row_top + Inches(0.35), Inches(6.5), Inches(0.35),
            font_size=Pt(9.5), color=C_GRAY)

    # sisi kanan: MQTT topics
    box(s, Inches(8.0), Inches(1.3), Inches(5.0), Inches(5.5),
        fill_color=C_NAVY)
    txt(s, "MQTT Topics  (HiveMQ Retained)",
        Inches(8.1), Inches(1.38), Inches(4.8), Inches(0.4),
        font_size=Pt(13), bold=True, color=C_ORANGE)
    accent_line(s, Inches(8.0), Inches(1.82), Inches(5.0), C_ORANGE, Pt(2))

    topics = [
        ("rollators/<id>/gas",    '{"gas": bool, "deviceMillis": int}'),
        ("rollators/<id>/sos",    '{"sos": bool, "deviceMillis": int}'),
        ("rollators/<id>/imu",    '{"status": "jalan|diam", "walking": bool,\n "pitch": float, "motionScore": float}'),
        ("rollators/<id>/status", '{"online": bool}  ← Last Will offline'),
    ]
    t_top = Inches(2.0)
    for topic, payload in topics:
        txt(s, topic, Inches(8.15), t_top, Inches(4.7), Inches(0.35),
            font_size=Pt(10.5), bold=True, color=C_CYAN)
        txt(s, payload, Inches(8.15), t_top + Inches(0.35), Inches(4.7), Inches(0.55),
            font_size=Pt(9), color=C_WHITE, italic=True)
        t_top += Inches(1.1)

    # catatan dashboard
    box(s, Inches(0.3), Inches(6.88), Inches(12.7), Inches(0.42),
        fill_color=RGBColor(0xE8,0xF8,0xF5))
    txt(s, "\u2139  dashboard.html — standalone HTML, polling /api/status setiap 1 detik. "
           "Buka via http://rorro.local atau IP lokal ESP32.",
        Inches(0.45), Inches(6.92), Inches(12.3), Inches(0.35),
        font_size=Pt(10), color=C_BLUE)


def slide_demo(prs):
    s = blank_slide(prs)
    bg(s, C_NAVY)
    section_header(s, "07", "Demo & Screenshot", "Demonstrasi sistem berjalan")
    footer(s, 8, 10)

    placeholders = [
        ("Hardware Rollator", "Foto fisik rollator + ESP32 terpasang"),
        ("Serial Monitor", "Log FreeRTOS: TaskMotor, TaskIMU, TaskSafety"),
        ("dashboard.html", "Browser membuka http://rorro.local"),
        ("Firebase Console", "Firestore — dokumen rollator + subcollection history"),
        ("Flutter App", "Tampilan realtime dari MQTT push"),
        ("MQTT Monitor", "mqtt_web_monitor.html — topic live feed"),
    ]

    cols = 3
    ph_w = Inches(4.1)
    ph_h = Inches(2.3)
    for i, (title, desc) in enumerate(placeholders):
        col = i % cols
        row = i // cols
        left = Inches(0.3) + col * Inches(4.35)
        top  = Inches(1.3) + row * Inches(2.6)
        box(s, left, top, ph_w, ph_h,
            fill_color=RGBColor(0x0D,0x22,0x36),
            line_color=C_CYAN, line_width=Pt(1.5))
        txt(s, "[FOTO / SCREENSHOT]",
            left, top + Inches(0.6), ph_w, Inches(0.55),
            font_size=Pt(13), color=RGBColor(0x55,0x65,0x78),
            align=PP_ALIGN.CENTER, italic=True)
        txt(s, title,
            left + Inches(0.1), top + Inches(1.35), ph_w - Inches(0.2), Inches(0.35),
            font_size=Pt(12), bold=True, color=C_CYAN)
        txt(s, desc,
            left + Inches(0.1), top + Inches(1.75), ph_w - Inches(0.2), Inches(0.4),
            font_size=Pt(9.5), color=C_WHITE)

    txt(s, "Ganti placeholder [FOTO / SCREENSHOT] dengan gambar aktual saat presentasi.",
        Inches(0.3), Inches(6.8), Inches(12.7), Inches(0.35),
        font_size=Pt(10), color=C_GRAY, align=PP_ALIGN.CENTER, italic=True)


def slide_tim(prs):
    s = blank_slide(prs)
    bg(s, C_LIGHT)
    section_header(s, "08", "Tim & Kontribusi", "Pengembang firmware + mobile")
    footer(s, 9, 10)

    # tim firmware
    box(s, Inches(0.3), Inches(1.3), Inches(6.0), Inches(5.5), fill_color=C_NAVY)
    txt(s, "\u26A1  Tim Firmware (ESP32)",
        Inches(0.45), Inches(1.38), Inches(5.7), Inches(0.45),
        font_size=Pt(16), bold=True, color=C_CYAN)
    accent_line(s, Inches(0.3), Inches(1.88), Inches(6.0), C_CYAN, Pt(2))

    fw_items = [
        ("Arsitektur FreeRTOS", "Dual-core task, portMUX locking"),
        ("Motor Control",       "BTS7960, PWM, potensiometer, triple-tap latch"),
        ("IMU Manager",         "MPU6050, motionScore EMA, state machine"),
        ("Safety Module",       "SOS, buzzer, priority 3 task"),
        ("Firebase Manager",    "Firestore REST, auth, deferred-queue"),
        ("MQTT Manager",        "HiveMQ TLS, retained publish, Last Will"),
        ("WiFi Manager",        "AP/STA state machine, mDNS, Preferences"),
        ("HTTP API",            "WebServer, dashboard.html, CORS"),
    ]
    item_top = Inches(2.05)
    for contrib, desc in fw_items:
        txt(s, f"\u25B6  {contrib}",
            Inches(0.45), item_top, Inches(2.6), Inches(0.38),
            font_size=Pt(11), bold=True, color=C_CYAN)
        txt(s, desc,
            Inches(3.1), item_top, Inches(3.1), Inches(0.38),
            font_size=Pt(10.5), color=C_WHITE)
        item_top += Inches(0.55)

    # tim mobile
    box(s, Inches(6.5), Inches(1.3), Inches(6.5), Inches(5.5),
        fill_color=RGBColor(0x1A,0x2A,0x40))
    txt(s, "\U0001F4F1  Tim Mobile (Flutter)",
        Inches(6.65), Inches(1.38), Inches(6.2), Inches(0.45),
        font_size=Pt(16), bold=True, color=C_ORANGE)
    accent_line(s, Inches(6.5), Inches(1.88), Inches(6.5), C_ORANGE, Pt(2))

    mobile_items = [
        ("Flutter App",         "Companion app untuk monitoring rollator"),
        ("MQTT Subscribe",      "rollators/<id>/# — realtime state push"),
        ("Dashboard Realtime",  "Gas, SOS, IMU walking status, pitch/roll"),
        ("History & Audit",     "Firebase Firestore subcollection query"),
        ("Offline Handling",    "status.online=false → IMU dianggap offline"),
        ("Notifikasi SOS",      "Push notification saat SOS aktif"),
    ]
    item_top = Inches(2.05)
    for contrib, desc in mobile_items:
        txt(s, f"\u25B6  {contrib}",
            Inches(6.65), item_top, Inches(2.5), Inches(0.38),
            font_size=Pt(11), bold=True, color=C_ORANGE)
        txt(s, desc,
            Inches(9.2), item_top, Inches(3.7), Inches(0.38),
            font_size=Pt(10.5), color=C_WHITE)
        item_top += Inches(0.7)

    txt(s, "Nama anggota tim dapat ditambahkan di sini sesuai kebutuhan presentasi.",
        Inches(0.3), Inches(6.88), Inches(12.7), Inches(0.35),
        font_size=Pt(9.5), color=C_GRAY, align=PP_ALIGN.CENTER, italic=True)


def slide_kesimpulan(prs):
    s = blank_slide(prs)
    bg(s, C_NAVY)
    section_header(s, "09", "Kesimpulan & Q&A", "")
    footer(s, 10, 10)

    # kiri: pencapaian
    box(s, Inches(0.3), Inches(1.3), Inches(6.2), Inches(5.5),
        fill_color=RGBColor(0x0D,0x22,0x36), line_color=C_GREEN, line_width=Pt(2))
    txt(s, "\u2714  Yang Berhasil Diimplementasikan",
        Inches(0.45), Inches(1.38), Inches(5.9), Inches(0.45),
        font_size=Pt(14), bold=True, color=C_GREEN)
    accent_line(s, Inches(0.3), Inches(1.88), Inches(6.2), C_GREEN, Pt(2))

    achieved = [
        "Dual-core FreeRTOS: TaskMotor, TaskIMU, TaskSafety",
        "Kendali motor DC via BTS7960 + potentiometer",
        "Deteksi berjalan/diam — MPU6050 + EMA motionScore",
        "Tombol SOS + buzzer state machine",
        "Deferred-queue: tasks tidak blocking networking",
        "Firebase Firestore — storage & history tiap event",
        "MQTT HiveMQ TLS — push realtime ke Flutter",
        "WiFi provisioning AP/STA + mDNS rorro.local",
        "HTTP API lokal + dashboard.html",
        "Tools Python: imu_tune.py, mqtt_test.py",
    ]
    a_top = Inches(2.0)
    for item in achieved:
        txt(s, f"\u2714  {item}", Inches(0.45), a_top, Inches(5.9), Inches(0.43),
            font_size=Pt(11), color=C_WHITE)
        a_top += Inches(0.45)

    # kanan: pengembangan selanjutnya
    box(s, Inches(6.7), Inches(1.3), Inches(6.3), Inches(2.8),
        fill_color=RGBColor(0x0D,0x22,0x36), line_color=C_ORANGE, line_width=Pt(2))
    txt(s, "\U0001F527  Pengembangan Selanjutnya",
        Inches(6.85), Inches(1.38), Inches(6.0), Inches(0.45),
        font_size=Pt(14), bold=True, color=C_ORANGE)
    accent_line(s, Inches(6.7), Inches(1.88), Inches(6.3), C_ORANGE, Pt(2))

    future = [
        "Encoder roda — odometri & feedback loop PID",
        "Monitoring baterai (GPIO 36 ADC)",
        "Fall detection — pitch/roll threshold otomatis",
        "TLS certificate verification (keamanan produksi)",
        "Physical emergency stop button",
    ]
    f_top = Inches(2.0)
    for item in future:
        txt(s, f"\u25B6  {item}", Inches(6.85), f_top, Inches(6.0), Inches(0.38),
            font_size=Pt(11), color=C_WHITE)
        f_top += Inches(0.43)

    # Q&A box
    box(s, Inches(6.7), Inches(4.35), Inches(6.3), Inches(2.45),
        fill_color=C_BLUE)
    txt(s, "Terima Kasih",
        Inches(6.7), Inches(4.5), Inches(6.3), Inches(0.65),
        font_size=Pt(30), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    txt(s, "RoRo — Smart Robot Rollator",
        Inches(6.7), Inches(5.15), Inches(6.3), Inches(0.4),
        font_size=Pt(14), color=C_CYAN, align=PP_ALIGN.CENTER)
    txt(s, "Ada pertanyaan?",
        Inches(6.7), Inches(5.6), Inches(6.3), Inches(0.4),
        font_size=Pt(14), color=C_WHITE, align=PP_ALIGN.CENTER, italic=True)
    txt(s, "rorro.local  \u2022  rollators/<id>/#",
        Inches(6.7), Inches(6.05), Inches(6.3), Inches(0.35),
        font_size=Pt(11), color=RGBColor(0xB0,0xC8,0xE0), align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main():
    try:
        from pptx import Presentation as _test  # noqa
    except ImportError:
        print("ERROR: python-pptx belum terinstall.")
        print("Jalankan:  pip install python-pptx")
        sys.exit(1)

    prs = prs_new()

    print("Membuat slide...")
    slide_cover(prs)
    print("  [1/10] Cover")
    slide_latar_belakang(prs)
    print("  [2/10] Latar Belakang")
    slide_arsitektur(prs)
    print("  [3/10] Arsitektur Sistem")
    slide_fitur_utama(prs)
    print("  [4/10] Fitur Utama")
    slide_tech_stack(prs)
    print("  [5/10] Tech Stack")
    slide_alur_kerja(prs)
    print("  [6/10] Alur Kerja")
    slide_api_dashboard(prs)
    print("  [7/10] API & Dashboard")
    slide_demo(prs)
    print("  [8/10] Demo & Screenshot")
    slide_tim(prs)
    print("  [9/10] Tim & Kontribusi")
    slide_kesimpulan(prs)
    print("  [10/10] Kesimpulan & Q&A")

    output = "RoRo_Presentasi.pptx"
    prs.save(output)
    print(f"\nSelesai! File: {output}")
    print("Buka dengan PowerPoint / LibreOffice Impress / Google Slides.")
    print("\nCatatan slide Demo (slide 8):")
    print("  Ganti placeholder [FOTO / SCREENSHOT] dengan gambar aktual.")


if __name__ == "__main__":
    main()
